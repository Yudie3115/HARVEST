"""
Processing pipeline – download, extract, classify, chunk, embed, and store documents
"""
import tempfile
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Optional
from uuid import UUID

import httpx
from loguru import logger

try:
    from prefect import flow, task
except ImportError:
    # The lightweight API image executes these helpers directly for Admin
    # uploads and intentionally does not ship the Prefect orchestration stack.
    def _direct_decorator(*decorator_args, **decorator_kwargs):
        if decorator_args and callable(decorator_args[0]) and len(decorator_args) == 1:
            return decorator_args[0]
        return lambda function: function

    flow = _direct_decorator
    task = _direct_decorator

from src.config.site_configs import get_site_name_for_url
from src.config.settings import settings
from src.llm.client import LLMClient
from src.llm.embeddings import EmbeddingClient
from src.storage.postgres import PostgresClient
from src.storage.minio_client import MinIOClient
from src.storage.vector_db import VectorDBClient
from src.storage.models import DocumentStatus
from src.utils.hashing import hash_content
from src.utils.text_processing import clean_text, chunk_text


# ---------------------------------------------------------------------------
# Helper: consistent, explanatory processing outcome logs
# ---------------------------------------------------------------------------


def _log_processing_status(
    *,
    document_id: str,
    url: str,
    status: str,
    explanation: str,
) -> None:
    """Log a document status with enough context to understand its outcome."""
    message = (
        f"Document processing status={status} document_id={document_id} "
        f"url={url}: {explanation}"
    )
    if status == "failed":
        logger.error(message)
    elif status == "skipped":
        logger.warning(message)
    else:
        logger.info(message)


# ---------------------------------------------------------------------------
# Helper: determine document type from URL / mime type
# ---------------------------------------------------------------------------


def _detect_document_type(url: str, mime_type: str = "") -> str:
    url_lower = url.lower()
    mime_lower = mime_type.lower()

    if url_lower.endswith(".pdf") or "pdf" in mime_lower:
        return "pdf"
    if any(url_lower.endswith(ext) for ext in (".mp4", ".avi", ".mov", ".webm", ".mkv")):
        return "video"
    if "youtube.com" in url_lower or "youtu.be" in url_lower or "vimeo.com" in url_lower:
        return "video"
    if any(url_lower.endswith(ext) for ext in (".docx", ".doc")):
        return "word"
    if any(url_lower.endswith(ext) for ext in (".xlsx", ".xls", ".csv")):
        return "spreadsheet"
    return "html"


# ---------------------------------------------------------------------------
# Tasks
# ---------------------------------------------------------------------------


def _load_uploaded_document(db: PostgresClient, document_id: str, url: str) -> Optional[Dict]:
    """Read a manual upload's stored bytes back from MinIO.

    upload:// URIs are synthetic, so retrying one (e.g. via reprocess-stuck)
    must reuse the file saved at ingestion time rather than fetch over HTTP.
    """
    document = db.get_document(UUID(document_id))
    file_path = document.file_path if document else None
    try:
        if not file_path:
            raise FileNotFoundError(
                "The uploaded file is no longer stored; upload it again to process it."
            )
        content = MinIOClient().get_object(file_path)
    except Exception as e:
        logger.error(f"Could not load uploaded file for {url}: {e}")
        db.update_document_status(UUID(document_id), DocumentStatus.FAILED, error_log=str(e))
        return None

    mime_type = document.mime_type or "application/octet-stream"
    db.update_document_status(UUID(document_id), DocumentStatus.QUEUED)
    return {
        "minio_path": file_path,
        "file_hash": hash_content(content),
        "file_size": len(content),
        "mime_type": mime_type,
        "doc_type": document.document_type or _detect_document_type(url, mime_type),
        "content": content,
    }


@task(
    name="download_document",
    retries=3,
    retry_delay_seconds=30,
)
def download_document(document_id: str, url: str) -> Optional[Dict]:
    """
    Download a document and store it in MinIO.

    Returns:
        Dict with keys: minio_path, file_hash, file_size, mime_type, doc_type
        or None on failure.
    """
    logger.info(f"Downloading: {url}")

    db = PostgresClient()
    if url.startswith("upload://"):
        return _load_uploaded_document(db, document_id, url)

    headers = {"User-Agent": settings.crawl_user_agent}
    site_name = get_site_name_for_url(url)

    try:
        with httpx.Client(
            headers=headers,
            timeout=settings.crawl_timeout_seconds,
            follow_redirects=True,
            verify=settings.verify_ssl_for_site(site_name),
        ) as client:
            response = client.get(url)
            response.raise_for_status()

        content = response.content
        mime_type = response.headers.get("content-type", "").split(";")[0].strip()
        doc_type = _detect_document_type(url, mime_type)
        file_hash = hash_content(content)
        file_size = len(content)

        # Store in MinIO
        minio = MinIOClient()
        suffix = Path(url.split("?")[0]).suffix or f".{doc_type}"
        object_name = f"{document_id}{suffix}"
        minio.upload_bytes(content, object_name=object_name, content_type=mime_type)

        # Update document record
        db.update_document_status(UUID(document_id), DocumentStatus.QUEUED)

        logger.info(f"Downloaded {url} → {object_name} " f"({file_size} bytes, type={doc_type})")
        return {
            "minio_path": object_name,
            "file_hash": file_hash,
            "file_size": file_size,
            "mime_type": mime_type,
            "doc_type": doc_type,
            "content": content,
        }

    except Exception as e:
        logger.error(f"Download failed for {url}: {e}")
        db.update_document_status(
            UUID(document_id),
            DocumentStatus.FAILED,
            error_log=str(e),
        )
        return None


@task(name="extract_content")
def extract_content(download_result: Dict, url: str = "") -> Optional[str]:
    """
    Extract plain text from downloaded content based on document type.

    Returns:
        Extracted text, or None on failure.
    """
    if not download_result:
        return None

    doc_type = download_result.get("doc_type", "html")
    content = download_result.get("content", b"")

    try:
        if doc_type == "pdf":
            with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
                tmp.write(content)
                tmp_path = tmp.name
            from src.extractors.pdf_extractor import extract_pdf_full

            result = extract_pdf_full(tmp_path)
            text = result.get("text", "")
            Path(tmp_path).unlink(missing_ok=True)

        elif doc_type == "word":
            from io import BytesIO
            from docx import Document as WordDocument

            document = WordDocument(BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
            text = "\n\n".join(parts)

        elif doc_type == "video":
            from src.extractors.video_extractor import extract_video_full

            video_data = extract_video_full(url)
            text = video_data.get("text", "")

        else:  # html or unknown
            from src.extractors.html_extractor import extract_html_content

            text = extract_html_content(content.decode("utf-8", errors="replace"), url=url) or ""

        text = clean_text(text)

        if len(text) < settings.min_content_length:
            logger.warning(f"Extracted text too short ({len(text)} chars) for {url}")
            return None

        # Truncate to max allowed
        if len(text) > settings.max_content_length:
            text = text[: settings.max_content_length]

        logger.info(f"Extracted {len(text)} chars from {url} (type={doc_type})")
        return text

    except Exception as e:
        logger.error(f"Content extraction failed for {url}: {e}")
        return None


@task(name="classify_document_content", retries=1, retry_delay_seconds=30)
def classify_document_content(text: str, url: str = "") -> Dict:
    """
    Use LLM to classify extracted content and build rich metadata.

    Returns:
        Classification dict.
    """
    if not text:
        return LLMClient._default_classification("")

    try:
        llm = LLMClient()
        classification = llm.classify_content(text, url=url)
        logger.debug(
            f"Classified content: relevance={classification.get('relevance_score', 0):.1f}"
        )
        return classification
    except Exception as e:
        logger.error(f"Content classification failed: {e}")
        return LLMClient._default_classification(text)


@task(name="chunk_and_embed", retries=1, retry_delay_seconds=30)
def chunk_and_embed(
    text: str,
    document_id: str,
    classification: Dict,
    url: str = "",
) -> List[Dict]:
    """
    Chunk the extracted text and generate an embedding for each chunk.

    Returns:
        List of dicts, each containing: text, vector, chunk_index, token_count, payload
    """
    if not text:
        return []

    chunks = chunk_text(text)
    if not chunks:
        return []

    embedder = EmbeddingClient()
    vectors = embedder.embed_batch(chunks)

    from src.utils.text_processing import estimate_tokens

    results = []
    for i, (chunk_text_val, vector) in enumerate(zip(chunks, vectors)):
        payload = {
            "document_id": document_id,
            "text": chunk_text_val,
            "topics": classification.get("topics", []),
            "crops": classification.get("crops", []),
            "regions": classification.get("regions", []),
            "source_url": url,
            "document_type": classification.get("document_type", "unknown"),
            "relevance_score": classification.get("relevance_score", 0),
            "created_at": datetime.utcnow().isoformat(),
        }
        results.append(
            {
                "chunk_index": i,
                "text": chunk_text_val,
                "vector": vector,
                "token_count": estimate_tokens(chunk_text_val),
                "payload": payload,
            }
        )

    logger.info(f"Created {len(results)} chunks with embeddings for document {document_id}")
    return results


@task(name="store_chunks")
def store_chunks(
    chunks: List[Dict],
    document_id: str,
    classification: Dict,
) -> int:
    """
    Persist chunks to PostgreSQL (ContentChunk table) and the configured vector DB.

    Returns:
        Number of chunks stored.
    """
    if not chunks:
        return 0

    db = PostgresClient()
    vector_db = VectorDBClient()
    vector_db.ensure_collection()

    chunk_ids: List[str] = []
    vectors: List[List[float]] = []
    payloads: List[Dict] = []

    for chunk in chunks:
        from uuid import uuid4

        chunk_uuid = str(uuid4())

        # PostgreSQL
        with db.get_session() as session:
            from src.storage.models import ContentChunk

            db_chunk = ContentChunk(
                document_id=UUID(document_id),
                chunk_index=chunk["chunk_index"],
                text=chunk["text"],
                token_count=chunk["token_count"],
                topics=classification.get("topics", []),
                crops=classification.get("crops", []),
                regions=classification.get("regions", []),
                vector_id=chunk_uuid,
                extra_metadata={
                    "relevance_score": classification.get("relevance_score", 0),
                    "document_type": classification.get("document_type"),
                },
            )
            session.add(db_chunk)
            session.commit()

        chunk_ids.append(chunk_uuid)
        vectors.append(chunk["vector"])
        payloads.append({**chunk["payload"], "chunk_id": chunk_uuid})

    # Vector DB bulk upsert
    if chunk_ids:
        vector_db.upsert_chunks(chunk_ids, vectors, payloads)

    logger.info(f"Stored {len(chunks)} chunks for document {document_id}")
    return len(chunks)


# ---------------------------------------------------------------------------
# Main processing flow
# ---------------------------------------------------------------------------


@flow(
    name="process_document",
    retries=1,
    retry_delay_seconds=120,
)
def process_document(document_id: str, url: str, site_name: str = "") -> Dict:
    """
    Full processing pipeline for a single document.

    Steps:
    1. Download & store in MinIO
    2. Extract text content
    3. LLM classification
    4. Chunk + embed
    5. Store chunks in PostgreSQL + vector DB
    6. Update document status in PostgreSQL

    Args:
        document_id: UUID string of the Document record
        url:         Source URL
        site_name:   Name of the source site

    Returns:
        Summary dict
    """
    db = PostgresClient()

    try:
        db.update_document_status(UUID(document_id), DocumentStatus.PROCESSING)
        _log_processing_status(
            document_id=document_id,
            url=url,
            status="processing",
            explanation="pipeline started; the document is being downloaded and indexed",
        )

        # 1. Download
        download_result = download_document(document_id, url)
        if not download_result:
            # download_document already logged the error and set status to FAILED
            _log_processing_status(
                document_id=document_id,
                url=url,
                status="failed",
                explanation=(
                    "download or raw-file storage failed; the document was not extracted "
                    "or indexed (see the preceding download error)"
                ),
            )
            return {
                "document_id": document_id,
                "url": url,
                "status": "failed",
                "chunks": 0,
                "reason": "download_failed",
            }

        # 2. Extract
        text = extract_content(download_result, url=url)
        if not text:
            db.update_document_status(
                UUID(document_id),
                DocumentStatus.SKIPPED,
                error_log="No extractable content",
            )
            _log_processing_status(
                document_id=document_id,
                url=url,
                status="skipped",
                explanation=(
                    "no usable text was extracted (content may be empty, too short, "
                    "unsupported, or image-only); no chunks were indexed"
                ),
            )
            return {
                "document_id": document_id,
                "url": url,
                "status": "skipped",
                "chunks": 0,
                "reason": "no_extractable_content",
            }

        # 3. Classify
        classification = classify_document_content(text, url=url)

        # Store classification
        db.create_classification(UUID(document_id), classification)

        # 4. Chunk + embed
        chunks = chunk_and_embed(text, document_id, classification, url=url)

        # 5. Persist chunks
        stored = store_chunks(chunks, document_id, classification)

        # 6. Mark as processed
        db.update_document_status(UUID(document_id), DocumentStatus.PROCESSED)

        summary = {
            "document_id": document_id,
            "url": url,
            "status": "processed",
            "chunks": stored,
            "relevance_score": classification.get("relevance_score", 0),
        }
        _log_processing_status(
            document_id=document_id,
            url=url,
            status="processed",
            explanation=(
                f"pipeline completed successfully; {stored} chunks were stored and indexed "
                f"(relevance_score={summary['relevance_score']})"
            ),
        )
        return summary

    except Exception as e:
        db.update_document_status(
            UUID(document_id),
            DocumentStatus.FAILED,
            error_log=str(e),
        )
        _log_processing_status(
            document_id=document_id,
            url=url,
            status="failed",
            explanation=(
                f"pipeline raised {type(e).__name__} during extraction, classification, "
                f"embedding, or storage; document was not fully indexed: {e}"
            ),
        )
        raise


@flow(name="process_documents_batch")
def process_documents_batch(
    documents: List[Dict],
    max_concurrent: int = 5,
) -> List[Dict]:
    """
    Process a batch of documents sequentially.

    In Prefect 2.x, @flow objects do not have .submit() – that API is
    task-only. Documents are therefore processed one at a time. For
    parallel processing, deploy flows to a work pool and run multiple
    workers (``prefect worker start --pool <pool> --limit <n>``).

    Args:
        documents:      List of dicts with keys: id, url, site_name
        max_concurrent: Reserved for future use (ignored in this version)

    Returns:
        List of per-document result dicts
    """
    logger.info(f"Processing batch of {len(documents)} documents")

    results = []
    for doc in documents:
        try:
            result = process_document(
                document_id=str(doc["id"]),
                url=doc["url"],
                site_name=doc.get("site_name", ""),
            )
            results.append(result)
        except Exception as e:
            results.append(
                {
                    "document_id": str(doc.get("id", "")),
                    "url": doc.get("url", ""),
                    "status": "failed",
                    "chunks": 0,
                    "reason": "processing_exception",
                    "error": str(e),
                }
            )

    status_counts = {
        status: sum(1 for result in results if result.get("status") == status)
        for status in ("processed", "skipped", "failed")
    }
    logger.info(
        f"Document processing batch complete: total={len(documents)}, "
        f"processed={status_counts['processed']} (successfully indexed), "
        f"skipped={status_counts['skipped']} (intentionally not indexed), "
        f"failed={status_counts['failed']} (pipeline error)"
    )
    return results
